"""
Exportación única y explícita. Por cada análisis, una carpeta:

  CARPETA/<analysis>/
  ├── objects.csv, image.csv, object_colors.csv...   una tabla por nivel (Image_ID + Image_name)
  ├── image.csv        una fila por foto: estado, mediciones de la foto, escala, áreas
  ├── metadata.csv     versión, fecha, análisis, parámetros, segmentación, correcciones
  ├── <analysis>.xlsx  lo mismo, una hoja por tabla
  └── views/<vista>/   imágenes de cada vista (opcional; se recalculan, ver core.batch)

Todo en inglés; CSV utf-8-sig. Image_ID = posición de la foto en la lista del proyecto,
igual en todos los análisis.
"""
from __future__ import annotations

import csv
import datetime as dt
import platform
from pathlib import Path
from typing import Any

import cv2
import numpy as np

from fenotit import APP_NAME, __version__, log

_log = log.get("export")

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    XLSX_OK = True
except ImportError:
    XLSX_OK = False

SKIP_PARAMS = ("roi_shapes", "mm_per_pixel", "color_space_code")


def save_image(folder: Path, name: str, img: np.ndarray):
    """Guarda imagen (rutas Unicode)."""
    folder.mkdir(parents=True, exist_ok=True)
    ext = Path(name).suffix.lower() or ".jpg"
    if ext not in (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp"):
        name, ext = Path(name).stem + ".jpg", ".jpg"
    ok, buf = cv2.imencode(ext, img)
    if not ok:
        raise OSError(f"no se pudo codificar {name}")
    (folder / name).write_bytes(buf.tobytes())


def _plain(v):
    if isinstance(v, np.generic):
        return v.item()
    if isinstance(v, (list, tuple, dict)):
        return str(v)
    return "" if v is None else v


def write_csv(path: Path, rows: list[dict]):
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = list(dict.fromkeys(k for r in rows for k in r))
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows({k: _plain(v) for k, v in r.items()} for r in rows)


def _params_of(result) -> dict:
    p = result.extra.get("params") or {}
    return {k: v for k, v in sorted(p.items()) if k not in SKIP_PARAMS}


def collect(analysis: str, results: dict[str, Any], images: list[str],
            image_info: dict[str, dict] | None = None, skipped: set | None = None,
            board: set | None = None, ids: dict[str, int] | None = None) -> dict[str, list[dict]]:
    """Tablas del análisis. `image` lleva una fila por foto de `images` (también las no
    analizadas): estado, sus mediciones de imagen y escala/áreas. `ids` = Image_ID de cada
    foto (su posición en el proyecto) cuando se exporta solo una parte."""
    image_info = image_info or {}
    skipped = skipped or set()
    tables: dict[str, list[dict]] = {"image": []}
    params = [_params_of(r) for r in results.values() if r.status == "ok"]
    varied = any(p != params[0] for p in params[1:]) if params else False
    for k, path in enumerate(images, 1):
        image_id = (ids or {}).get(path, k)
        name = Path(path).name
        r = results.get(path)
        status = ("calibration" if path in (board or ()) else "skipped" if path in skipped else "not_analyzed") \
            if r is None else \
            ("ok" if r.status == "ok" else "error")
        head = {"Image_ID": image_id, "Image_name": name, "status": status}
        tail = dict(image_info.get(path, {}))
        if r is not None and r.status != "ok":
            tail["error"] = r.error
        if varied and r is not None and r.status == "ok":
            tail["params"] = "; ".join(f"{k}={v}" for k, v in _params_of(r).items())
        own = (r.extra.get("tables") or {}).get("image") if r is not None and r.status == "ok" else None
        tables["image"].extend({**head, **x, **tail} for x in (own or [{}]))
        if r is None or r.status != "ok":
            continue
        for table, rows in (r.extra.get("tables") or {}).items():
            if table != "image":
                tables.setdefault(table, []).extend(
                    {"Image_ID": image_id, "Image_name": name, **x} for x in rows)
    return tables


def metadata(analysis: str, results: dict[str, Any], project=None, extra: dict | None = None) -> list[dict]:
    """Filas clave/valor: lo necesario para repetir y citar el análisis."""
    ok = [r for r in results.values() if r.status == "ok"]
    rows = [("software", APP_NAME), ("version", __version__),
            ("exported", dt.datetime.now().isoformat(timespec="seconds")),
            ("python", platform.python_version()), ("system", f"{platform.system()} {platform.release()}"),
            ("analysis", analysis), ("n_images_ok", len(ok)),
            ("n_images_error", sum(1 for r in results.values() if r.status != "ok"))]
    if project is not None:
        seg = project.segmentation
        corr = project.corrections
        rows += [("project", str(project.file or "")), ("length_unit", project.unit),
                 ("scale_all_images_mm_per_px", project.scale.mm_per_pixel or ""),
                 ("scale_source", project.scale.source),
                 ("images_with_own_scale", len(project.image_scale)),
                 ("analysis_areas_all_images", len(project.roi)),
                 ("images_with_own_areas", len(project.image_roi)),
                 ("objects_excluded_manually", sum(len(v) for v in getattr(project, "excluded", {}).values())),
                 ("segmentation", f"{seg.color_space} channel {seg.channel} "
                                  + ("Otsu" if seg.auto else f"{seg.min_val}-{seg.max_val}")),
                 ("corrections", ", ".join(corr.active()) or "none"),
                 ("max_megapixels", project.max_mpx or "full"),
                 ("preset", "none")]
    if ok:
        params = [_params_of(r) for r in ok]
        if any(p != params[0] for p in params[1:]):
            rows.append(("params", "varied between images: see the image table"))
        else:
            rows += [(f"param.{k}", v) for k, v in params[0].items()]
    rows += list((extra or {}).items())
    return [{"key": k, "value": _plain(v)} for k, v in rows]


def write_tables(folder: Path, tables: dict[str, list[dict]], meta: list[dict], excel: bool = True,
                 excel_name: str = "results.xlsx") -> list[Path]:
    """Un CSV por tabla, metadata.csv y (opcional) un Excel con una hoja por tabla."""
    folder.mkdir(parents=True, exist_ok=True)
    written = []
    for table, rows in tables.items():
        if rows:
            write_csv(folder / f"{table}.csv", rows)
            written.append(folder / f"{table}.csv")
    write_csv(folder / "metadata.csv", meta)
    written.append(folder / "metadata.csv")
    if excel and XLSX_OK:
        path = folder / excel_name
        _excel(path, {**{k: v for k, v in tables.items() if v}, "metadata": meta})
        written.append(path)
    return written


def _excel(path: Path, tables: dict[str, list[dict]]):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="2166AC")
    head_font = Font(color="FFFFFF", bold=True, size=9)
    cell_font = Font(size=9)
    border = Border(*(Side(style="thin", color="CCCCCC"),) * 4)
    for table, rows in tables.items():
        ws = wb.create_sheet(table[:31])
        headers = list(dict.fromkeys(k for r in rows for k in r))
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font, cell.fill, cell.border = head_font, head_fill, border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r_i, row in enumerate(rows, 2):
            for c, h in enumerate(headers, 1):
                val = _plain(row.get(h, ""))
                cell = ws.cell(row=r_i, column=c, value=val)
                cell.font, cell.border = cell_font, border
                cell.alignment = Alignment(horizontal="right" if isinstance(val, (int, float)) else "left")
        for c, h in enumerate(headers, 1):
            width = max([len(str(h))] + [len(str(_plain(r.get(h, "")))) for r in rows[:500]])
            ws.column_dimensions[get_column_letter(c)].width = min(width + 2, 40 if table == "metadata" else 25)
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def save_views(result, folder: Path, image_name: str, display: dict) -> int:
    """Imágenes de cada vista (con marcas y leyenda según `display`) en folder/<vista>/."""
    from fenotit.core.pipeline.views import decorate
    draw = decorate(result, display, Path(image_name).stem)
    n = 0
    for view, sub in (result.extra.get("step_folders") or {}).items():
        img = result.step_images.get(view)
        if img is not None:
            save_image(folder / sub, image_name, draw(view, img))
            n += 1
    return n
