"""Tabla de datos del usuario (CSV o Excel) que relaciona cada foto con genotipo,
repetición, tratamiento... Sus columnas sirven para agrupar los análisis."""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp"}


def read_table(path: str | Path) -> tuple[list[str], list[dict]]:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        cols = [str(c).strip() if c is not None else f"col{i + 1}" for i, c in enumerate(next(it))]
        rows = [dict(zip(cols, r)) for r in it if any(v not in (None, "") for v in r)]
        return cols, rows
    text = path.read_text(encoding="utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(text.splitlines(), dialect=dialect)
    cols = [c.strip() for c in reader.fieldnames or []]
    rows = [{c.strip(): (v.strip() if isinstance(v, str) else v) for c, v in r.items() if c}
            for r in reader]
    return cols, [r for r in rows if any(v not in (None, "") for v in r.values())]


def key(value) -> str:
    """Normaliza un nombre para comparar: sin extensión, sin mayúsculas, 001 = 1 = 1.0."""
    s = str(value).strip()
    if Path(s).suffix.lower() in IMAGE_EXTS:
        s = Path(s).stem
    s = s.casefold()
    try:
        f = float(s)
        if f.is_integer():
            return str(int(f))
    except ValueError:
        pass
    return s


@dataclass
class Match:
    by_image: dict[str, dict] = field(default_factory=dict)   # nombre de imagen -> fila
    images_without_row: list[str] = field(default_factory=list)
    rows_without_image: list[str] = field(default_factory=list)
    duplicated_keys: list[str] = field(default_factory=list)


def match(image_names: list[str], rows: list[dict], key_column: str) -> Match:
    index, dup = {}, []
    for r in rows:
        k = key(r.get(key_column, ""))
        if k in index:
            dup.append(str(r.get(key_column)))
        index.setdefault(k, r)
    m = Match(duplicated_keys=dup)
    used = set()
    for name in image_names:
        k = key(name)
        if k in index:
            m.by_image[name] = index[k]
            used.add(k)
        else:
            m.images_without_row.append(name)
    m.rows_without_image = [str(r.get(key_column)) for r in rows if key(r.get(key_column, "")) not in used]
    return m


def guess_key_column(columns: list[str], rows: list[dict], image_names: list[str]) -> str | None:
    """La columna que coincide con más nombres de imagen distintos."""
    names = {key(n) for n in image_names}
    best, best_n = None, 0
    for c in columns:
        n = len({key(r.get(c, "")) for r in rows} & names)
        if n > best_n:
            best, best_n = c, n
    return best
